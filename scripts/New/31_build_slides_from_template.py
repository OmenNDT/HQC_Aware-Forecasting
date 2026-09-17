#!/usr/bin/env python3
"""31_build_slides_from_template.py — Dựng slide báo cáo cuối từ mẫu '5. Mẫu Slide báo cáo.pptx' và nội dung Bao-cao-do-an-final.md.
Giữ master/theme của mẫu; giữ slide tiêu đề, mục lục và Hỏi–đáp của mẫu; các slide nội dung thêm bằng layout OBJECT / Hình ảnh."""
import copy, os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from PIL import Image

ROOT = "/home/sontn/Projects/HQC_Aware-Forecasting"; BC = f"{ROOT}/Bao_cao"
TPL = "/tmp/claude-1000/-home-sontn-Projects-HQC-Aware-Forecasting/455b4af4-2d17-44fb-87ca-effc26e4631c/scratchpad/slide-working.pptx"   # mẫu đã rút còn slide 0,1,11 bằng rearrange.py của skill pptx
OUT = f"{BC}/Slide-bao-cao-do-an-final.pptx"
prs = Presentation(TPL); L = {l.name: l for l in prs.slide_layouts}
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def set_text(tf, items, size=16, sub_size=14):
    """items: list of str hoặc (str, level). Xóa nội dung cũ, giữ định dạng layout."""
    tf.clear(); first = True
    for it in items:
        txt, lvl = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.level = lvl; r = p.add_run(); r.text = txt; r.font.size = Pt(size if lvl == 0 else sub_size)
        if txt.endswith(":") and lvl == 0: r.font.bold = True


def fit_picture(slide, path, rect):
    """Chèn ảnh vừa khít trong hộp (left, top, width, height) EMU, giữ tỉ lệ, căn giữa."""
    left, top, w, h = rect; iw, ih = Image.open(path).size; s = min(w / iw, h / ih); pw, ph = int(iw * s), int(ih * s)
    return slide.shapes.add_picture(path, left + (w - pw) // 2, top + (h - ph) // 2, pw, ph)


def add_object(title, items, size=16):
    s = prs.slides.add_slide(L["OBJECT"]); s.shapes.title.text = title
    body = [p for p in s.placeholders if p.placeholder_format.idx == 1][0]; set_text(body.text_frame, items, size); return s


def add_image_slide(title, items, image, size=14):
    s = prs.slides.add_slide(L["Hình ảnh"]); s.shapes.title.text = title
    body = [p for p in s.placeholders if p.placeholder_format.idx == 1][0]; set_text(body.text_frame, items, size, 12)
    pic = [p for p in s.placeholders if p.placeholder_format.idx == 2][0]; rect = (pic.left, pic.top, pic.width, pic.height)
    pic._element.getparent().remove(pic._element); fit_picture(s, image, rect); return s


def add_table(slide, rows, rect, col_widths=None, size=13, header_size=14):
    left, top, w, h = rect; nr, nc = len(rows), len(rows[0])
    tbl = slide.shapes.add_table(nr, nc, left, top, w, h).table
    if col_widths:
        tot = sum(col_widths)
        for j, cw in enumerate(col_widths): tbl.columns[j].width = int(w * cw / tot)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i, j); c.text = str(val)
            for p in c.text_frame.paragraphs:
                for r in p.runs: r.font.size = Pt(header_size if i == 0 else size); r.font.bold = (i == 0)
    return tbl


def clear_body(slide):
    for p in list(slide.placeholders):
        if p.placeholder_format.idx == 1: p._element.getparent().remove(p._element)


# ---------- mẫu đã rút còn 3 slide: tiêu đề, mục lục, hỏi đáp ----------
sldIdLst = prs.slides._sldIdLst
title, toc, qa = prs.slides[0], prs.slides[1], prs.slides[2]
for sh in title.shapes:
    if sh.is_placeholder and sh.placeholder_format.idx == 2:
        set_text(sh.text_frame, ["CẢNH BÁO SỚM BẤT THƯỜNG THIẾT BỊ QUAY BẰNG HỌC SÂU TÁI TẠO", "trên dữ liệu giám sát rung công nghiệp, triển khai trên ESP32-S3"], 24, 18)
    elif sh.has_text_frame and "GVHD" in sh.text_frame.text:
        set_text(sh.text_frame, ["GVHD: ThS. Phan Đình Duy", "Trương Ngọc Sơn – 25210183", "Trần Tín Nghĩa – 25210147", "Hồ Thị Mỹ Phương – 25210170"], 18, 18)
for sh in toc.shapes:
    if sh.is_placeholder and sh.placeholder_format.idx == 1:
        set_text(sh.text_frame, ["TỔNG QUAN ĐỀ TÀI", "CƠ SỞ LÝ THUYẾT", "PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG", "KẾT QUẢ THỰC NGHIỆM", "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN"], 24)

# ---------- 1. Tổng quan ----------
add_image_slide("1. TỔNG QUAN ĐỀ TÀI", [
    "Bài toán:", ("Bảo trì theo ngưỡng cố định: phát hiện muộn, không có dữ liệu hỏng gán nhãn, tải lẫn với hư hỏng", 1),
    ("Câu hỏi khác: máy còn giống lúc vận hành ổn định không, và đang rời xa nhanh đến đâu?", 1),
    "Đối tượng:", ("Bơm cấp nước lò hơi P29201A, Nhà máy Đạm Cà Mau; 4 ổ đỡ × 2 đầu đo (X, Y) = 8 kênh rung 1X (biên độ, pha) và Direct", 1),
    ("Ngưỡng nhà máy 65 µm peak-to-peak; hai lần dừng tháo kiểm tra 04/01 và 29/05/2026 để đo lead-time", 1),
    "Tham chiếu: Gribbestad và cộng sự, Entropy 2021 [1]: học sâu tái tạo, train chỉ trên dữ liệu bình thường",
], f"{BC}/P29201A Overview.jpg")
add_object("1.2 Mục tiêu", [
    "1. Kho đặc trưng 1X 8 kênh, bước 10 phút, trên toàn chuỗi 28/09/2025 → 02/09/2026",
    "2. Chọn nền gốc bằng thước ổn định của cả hệ 8 kênh; chia dữ liệu theo vai trò, không tham số nào chọn trên dữ liệu đánh giá",
    "3. Bảy họ mô hình tái tạo (PCA, AE, SAE, VAE, DBN, ED-LSTM, ED-CNN) qua cùng giao thức kiểm tra chéo theo khối",
    "4. Luật theo tốc độ thay đổi, ngưỡng chỉnh trên đoạn âm và khóa trước khi kiểm mù; ba tầng cảnh báo; đo lead-time trước hai mốc dừng thật",
    "5. Lớp giải thích: đóng góp từng kênh chỉ ra ổ đỡ đang suy giảm, dịch sang cơ chế cơ học theo ISO 13373",
    "6. Lượng tử hóa sang số nguyên, viết lại ba tầng bằng C, kiểm chứng cùng ngày báo với Python, biên dịch cho ESP32-S3",
], 16)

# ---------- 2. Cơ sở lý thuyết ----------
s = prs.slides.add_slide(L["So sanh 1"]); s.shapes.title.text = "2. CƠ SỞ LÝ THUYẾT"
ph = {p.placeholder_format.idx: p for p in s.placeholders}
set_text(ph[1].text_frame, ["Học sâu tái tạo [1]:", ("Huấn luyện chỉ trên dữ liệu bình thường, không nhãn", 1), ("Sai số tái tạo SPE = Σ(x − x̂)² làm chỉ số bất thường", 1),
    ("Đóng góp từng cảm biến để giải thích", 1), ("Sáu kiến trúc: AE, SAE, VAE, DBN, ED-LSTM, ED-CNN", 1),
    "Ba điểm yếu [1] tự nhận:", ("Ngưỡng 40/60 và sigmoid chỉnh tay", 1), ("Không có hàm chấm điểm để chọn kiến trúc", 1), ("Chỉ trả lời “có gì đó sai”, không trả lời “sai ở đâu” và “khi nào hỏng”", 1)], 15, 13)
set_text(ph[2].text_frame, ["Chẩn đoán rung theo bậc tốc độ quay (ISO 13373 [4]):", ("0,5X: xoáy dầu (oil whirl)", 1), ("1X: mất cân bằng; pha giữ → điểm nặng cố định, pha quay dần → điểm nặng di chuyển", 1),
    ("2X: lệch trục", 1), ("Tần số không đồng bộ dải rộng: hỏng vòng bi", 1),
    "Giới hạn kiểm soát thống kê:", ("Giới hạn 99% của SPE ngoài mẫu (Jackson–Mudholkar [2])", 1), ("Ngưỡng tuyệt đối theo ISO 20816 [3] là hướng mở rộng", 1)], 15, 13)

# ---------- 3. Thiết kế ----------
add_image_slide("3. THIẾT KẾ – 3.1 Kiến trúc", [
    "Nguồn đo: System 1 (1X Amp/Phase, Direct) → PI Data Archive",
    "Tầng dữ liệu: cổng lọc máy chạy, thống nhất tham chiếu pha, kho đặc trưng 24 chiều",
    "Tầng mô hình: nền gốc 9b → 5 khối chéo → bảy họ → SAE 3.600 tham số",
    "Ba tầng cảnh báo: A ngoại suy tới 65 µm; B tốc độ lệch khỏi nền; C chữ ký hình dạng hệ",
    "Thiết bị biên: SAE int16/int8 + ba tầng bằng C99 trên ESP32-S3, mỗi 10 phút",
], f"{BC}/hinh-kien-truc-tong-the.png", 14)
add_image_slide("3.2 Dữ liệu", [
    "42.064 mẫu 10 phút sau cổng lọc, 28/09/2025 → 02/09/2026; 7.129 giờ Direct",
    "Vai trò từng đoạn:", ("Đoạn 7, 8, 9a: đoạn âm, hệ không được báo", 1), ("9b (20/02–22/03/2026): nền gốc, huấn luyện", 1),
    ("9c: đợt leo tới dừng 2; nửa đầu chỉnh ngưỡng, nửa sau kiểm trước", 1), ("Đoạn 11: kiểm mù, chạy một lần sau khi khóa", 1),
    "Chọn 9b: dốc biên độ 2,0 %/tháng thấp nhất, trôi pha 4,4 °/tháng, mức 2001X 17,5 µm thấp nhất",
], f"{BC}/hinh-chon-nen-goc-9b.png", 13)
add_object("3.3 Đặc trưng và mô hình", [
    "Đầu vào 24 chiều: 8 biên độ 1X chuẩn hóa theo nền + 8 sin + 8 cos của pha",
    "Bảy họ mô hình, 27 cấu hình (mô hình ngẫu nhiên: 3 lần khởi tạo); không loại kiến trúc nào trước bằng lý lẽ",
    "Giao thức năm khối ba vai: nền 9b chia 5 khối ≈ 6 ngày, đệm 12 giờ; mỗi fold 1 khối kiểm tra, 1 khối dừng sớm, 3 khối huấn luyện",
    "Bốn tiêu chí chọn:", ("Ổn định giữa các khối nền: báo giả trung bình ≤ 3 %, lớn nhất ≤ 10 %", 1), ("Tách được: tỉ số SPE đợt leo / nền ≥ 1,3", 1),
    ("Lead-time đo được thật trước mốc dừng, không tính ngày báo bị giới hạn", 1), ("Giải thích đúng ổ đang thay đổi", 1),
    "Sai số tái tạo nhỏ nhất không phải tiêu chí",
], 15)
s = add_object("3.4 Ba tầng cảnh báo", []); clear_body(s)
add_table(s, [["Tầng", "Trục dữ liệu", "Cách tính", "Quy tắc báo"],
    ["A · Ngoại suy tới ngưỡng", "Direct 8 kênh", "Hồi quy 14 ngày chạy, ngoại suy tới 65 µm, cận dưới 90 %", "Kênh leo > 1,10 × trung vị 90 ngày, dốc có ý nghĩa, còn < 30 ngày, giữ 2 ngày"],
    ["B · Tốc độ lệch khỏi nền", "SPE so nền gốc 9b", "Dốc log SPE 14 ngày và 45 ngày", "τ₁₄ = 0,135 (3 ngày liên tiếp); τ₄₅ = 0,035 (5 ngày)"],
    ["C · Chữ ký hình dạng hệ", "Pha 1X 8 kênh theo tuần", "Trung bình tròn, tốc độ quay pha 4 tuần", "Đổi chiều quay ổ 2001X so 8 tuần trước, sd < 5°, giữ 2 tuần → “cần xem”"]],
    (Emu(774145), Emu(1400000), Emu(10579654), Emu(3000000)), [2.2, 2.0, 3.2, 4.2], 12, 13)
tb = s.shapes.add_textbox(Emu(774145), Emu(4600000), Emu(10579654), Emu(1200000)); set_text(tb.text_frame, [
    "Lead-time = ngày tầng báo lần đầu → ngày nhà máy dừng; đo riêng từng tầng, cho cả hai lần dừng",
    "Ngưỡng chỉnh trên đoạn âm và nửa đầu 9c, khóa kèm mã kiểm tra; đoạn 11 kiểm mù chạy đúng một lần"], 14)
add_object("3.5 Thiết kế nhúng", [
    "Thiết bị: ESP32-S3 DevKitC (2 lõi 240 MHz, 512 KB RAM, 16 MB flash)",
    "Mã C99 thuần, không thư viện, không cấp phát động; dùng chung cho bộ mô phỏng máy tính và firmware",
    "Không huấn luyện lại, không đổi ngưỡng; mọi hằng số sinh tự động từ file đã khóa kèm mã kiểm tra",
    "Chỉ tiêu chấp nhận:", ("Ngày báo B₁₄/B₄₅/A/C trùng 100 % với Python trên toàn chuỗi phát lại", 1),
    ("SPE lệch ≤ 5 % ở vùng có ý nghĩa; ở nền chỉ cần không đổi ngày báo", 1), ("< 32 KB flash hằng số, < 16 KB RAM trạng thái", 1),
    "Sơ đồ lượng tử hóa chọn bằng phép đo: phương án nào làm mất một ngày báo là loại",
], 15)

# ---------- 4. Kết quả ----------
add_image_slide("4. KẾT QUẢ – 4.1 Bảy họ", [
    "Ngưỡng tốc độ chung τ = 0,135/ngày chỉnh trên 16/27 cấu hình đạt tiêu chí 1–2",
    "Không cấu hình nào báo trên đoạn 7, 8, 9a",
    "SAE 3.600 tham số: báo 06/04, 53 ngày trước dừng 2; ổn định 0,01–0,02 / 0,05; tách 1.140–1.296; ổ 2001 chiếm 97 % đóng góp",
    "PCA k = 8 báo 25/03 nhưng bị giới hạn bởi ngày sớm nhất có thể báo; PCA k = 6 không đạt tiêu chí ổn định (27 %)",
    "Kiểm mù đoạn 11: 0 báo giả; đợt leo chậm 19/07 không mô hình nào báo (dốc 0,035–0,045 < τ)",
    "Quyết định: tầng B dùng SAE; thêm thang 45 ngày",
], f"{BC}/hinh-so-sanh-7-mo-hinh.png", 13)
add_image_slide("4.2 Ba tầng cảnh báo", [
    "Dừng 1 (04/01/2026): không tầng nào báo, đúng kỳ vọng ghi trước",
    "Dừng 2 (29/05/2026):", ("B₁₄: 06/04 → 53 ngày", 1), ("B₄₅: 11/04 → 48 ngày, giữ báo 25/29 ngày tháng 5", 1),
    ("A: 29/04 → 30 ngày, do ổ 2003Y, hợp lệ theo luật", 1), ("C: không báo (một tuần có dấu hiệu)", 1),
    "Đoạn âm: 0 báo ở A, B₁₄, B₄₅; C 1 tuần trong 21 ngày sau tháo lắp",
    "Kiểm mù đoạn 11: 0 báo giả; đợt leo chậm chưa tầng nào báo; A chưa kiểm được (Direct tới 21/07)",
], f"{BC}/hinh-ba-tang-toan-chuoi.png", 13)
s = add_object("4.3 Đối chiếu chỉ tiêu", []); clear_body(s)
add_table(s, [["Chỉ tiêu", "Mong đợi", "Kết quả"],
    ["Ổn định giữa các khối nền", "báo giả TB ≤ 3 %, lớn nhất ≤ 10 %", "SAE 0,01–0,02 / 0,05; 16/27 cấu hình đạt"],
    ["Không báo trên đoạn âm", "0", "0 ở cả bốn tầng (C: 1 tuần sau tháo lắp, tách riêng)"],
    ["Lead-time trước quyết định dừng", "đo được", "53 ngày (B₁₄), 48 (B₄₅), 30 (A) trước 29/05; dừng 1 không báo"],
    ["Kiểm mù: bình nguyên / leo chậm", "0 báo giả / bắt được", "0 báo giả / chưa bắt được"],
    ["Lớp giải thích", "chỉ đúng ổ", "ổ 2001 chiếm 97 % đóng góp"],
    ["Nhúng: tương đương Python", "ngày báo trùng 100 %", "trùng 100 % cả bốn tầng"],
    ["Nhúng: dung lượng", "< 32 KB flash, < 16 KB RAM", "9,9 KB flash, 8,3 KB RAM"]],
    (Emu(774145), Emu(1300000), Emu(10579654), Emu(4300000)), [3.0, 3.0, 5.0], 12, 13)
s = add_object("4.4 Trình diễn: phát lại theo ngày và video", []); clear_body(s)
fit_picture(s, f"{BC}/demo/khung/03-B14-bao-06-04.png", (Emu(774145), Emu(1300000), Emu(5200000), Emu(2925000)))
fit_picture(s, f"{BC}/demo/khung/05-tang-A-bao-29-04-2003Y.png", (Emu(6153800), Emu(1300000), Emu(5200000), Emu(2925000)))
tb = s.shapes.add_textbox(Emu(774145), Emu(4350000), Emu(10579654), Emu(1700000)); set_text(tb.text_frame, [
    "Trang phát lại tự chứa: đứng ở một ngày, chỉ thấy dữ liệu tới ngày đó; bốn đèn A / B₁₄ / B₄₅ / C theo bảng kết quả đã khóa",
    "Trái: 06/04, B₁₄ báo khi 2001X còn 25 µm Direct, SPE dưới giới hạn kiểm soát → báo bằng tốc độ, không bằng mức",
    "Phải: 29/04, tầng A báo do 2003Y (24 ngày tới 65). Video câm 6 phút quay màn hình thật, 13 mốc chú thích"], 13)
s = add_object("4.5 Lượng tử hóa và nhúng trên ESP32-S3", []); clear_body(s)
add_table(s, [["Phương án", "Trọng số", "Sai số SPE nền (p95)", "Ngày báo B₁₄ / B₄₅", "Kết luận"],
    ["Kích hoạt int8", "—", "83 %", "—", "loại"], ["Trọng số int8 toàn bộ", "3.456 B", "16 %", "mất 06/04 / giữ", "loại"],
    ["Lớp 0 int16, còn lại int8", "4.032 B", "6,4 %", "mất 06/04 / giữ", "loại"], ["Lớp 0–1 int16, lớp 2–5 int8", "4.608 B", "5,1 %", "giữ nguyên 19 / 49 ngày", "chọn"],
    ["int16 toàn bộ", "6.912 B", "2,1 %", "giữ / giữ", "không cần"]],
    (Emu(774145), Emu(1300000), Emu(10579654), Emu(2300000)), [3.2, 1.6, 2.2, 2.6, 1.4], 12, 13)
tb = s.shapes.add_textbox(Emu(774145), Emu(3750000), Emu(10579654), Emu(2300000)); set_text(tb.text_frame, [
    "Mã C ba tầng đối chiếu Python trên 42.064 mẫu: ngày báo B₁₄, B₄₅ trùng 100 %; tầng A 231/231 hàng trùng; tầng C 47/47 tuần trùng",
    "Firmware ESP-IDF 5.3.2: ảnh 294 KB; hằng số mô hình 9,9 KB flash; trạng thái ba tầng 8,3 KB RAM; 3.456 phép nhân-cộng mỗi 10 phút",
    "Rà soát và kiểm thử độc lập: suy luận số nguyên trùng từng bit với mô phỏng; ba lỗi mức cao sửa trước khi chốt",
    "Chờ bo: đo độ trễ, bộ nhớ động, dòng điện trên thiết bị thật"], 14)

# ---------- 5. Kết luận ----------
add_object("5. KẾT LUẬN", [
    "Đã đạt, có số đo:", ("Lead-time 53 và 48 ngày (tầng B), 30 ngày (tầng A) trước quyết định dừng thật; dừng 1 không báo, đúng kỳ vọng", 1),
    ("Không báo giả trên ba đoạn âm và bình nguyên kiểm mù; tham số khóa trước khi chấm", 1),
    ("Mô hình chọn bằng thực nghiệm có kiểm soát: 27 cấu hình, bảy họ, năm khối chéo → SAE", 1),
    ("Nhúng được kiểm chứng tương đương: 100 % ngày báo; firmware ESP32-S3 294 KB, 8,3 KB RAM", 1),
    "Giới hạn:", ("Đợt leo chậm từ 19/07 chưa tầng nào báo; tầng A chưa kiểm được vì Direct tới 21/07", 1),
    ("Một máy, hai mốc dừng tháo kiểm tra; lead-time đo tới quyết định vận hành, không tới hư hỏng xác nhận", 1),
    ("Firmware chưa lưu trạng thái qua mất điện; chưa có số đo trên bo", 1),
], 15)
add_object("5.2 Hướng phát triển", [
    "Nạp firmware lên ESP32-S3, phát lại 42.064 mẫu qua USB, đo độ trễ, bộ nhớ, dòng điện",
    "Kéo Direct từ 22/07 và 1X sau 02/09 để kiểm tầng A và luật C trên đợt leo chậm",
    "Lưu trạng thái ba tầng vào NVS, gốc tuần cố định, đóng ngày theo đồng hồ",
    "Bắt đợt leo chậm: thang 90 ngày hoặc luật mức theo SPE, chỉnh và khóa trên dữ liệu mới",
    "Tầng ngưỡng tuyệt đối theo ISO 20816 song song tầng thay đổi",
    "Kết nối thật qua Modbus TCP / OPC UA; mở rộng sang cặp bơm song song; đối chứng TensorFlow Lite Micro",
], 16)
add_object("Tài liệu tham khảo", [
    "[1] M. Gribbestad, M. U. Hassan, I. A. Hameed, K. Sundli, “Health Monitoring of Air Compressors Using Reconstruction-Based Deep Learning for Anomaly Detection with Increased Transparency,” Entropy, vol. 23, no. 1, p. 83, 2021.",
    "[2] J. E. Jackson, G. S. Mudholkar, “Control Procedures for Residuals Associated with Principal Component Analysis,” Technometrics, vol. 21, no. 3, pp. 341–349, 1979.",
    "[3] ISO 20816-1:2016, Mechanical vibration — Measurement and evaluation of machine vibration — Part 1: General guidelines.",
    "[4] ISO 13373-3:2015, Condition monitoring and diagnostics of machines — Vibration condition monitoring — Part 3: Guidelines for vibration diagnosis.",
], 14)
# đưa slide Hỏi–đáp xuống cuối
ids = list(sldIdLst); qa_id = ids[2]; sldIdLst.remove(qa_id); sldIdLst.append(qa_id)
prs.save(OUT); print("saved", OUT, len(prs.slides), "slides")
